from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
import os
import datetime

def create_presentation():
    # Create presentation
    prs = Presentation()

    # Define slide layouts
    title_slide_layout = prs.slide_layouts[0]
    bullet_slide_layout = prs.slide_layouts[1]
    
    # --- Slide 1: Title Slide ---
    slide1 = prs.slides.add_slide(title_slide_layout)
    title1 = slide1.shapes.title
    subtitle1 = slide1.placeholders[1]
    title1.text = "RentWheels"
    
    date_str = datetime.date.today().strftime("%B %d, %Y")
    subtitle1.text = f"MCA Final Year Project Presentation\nPremium Vehicle Rental Platform\n\n[Your Name] | {date_str}"

    # --- Slide 2: Problem Statement ---
    slide2 = prs.slides.add_slide(bullet_slide_layout)
    shapes2 = slide2.shapes
    title2 = shapes2.title
    body2 = shapes2.placeholders[1]
    title2.text = "The Problem Statement"
    tf2 = body2.text_frame
    
    p = tf2.paragraphs[0]
    p.text = "Current vehicle rental processes are highly fragmented and outdated."
    p.level = 0
    p = tf2.add_paragraph()
    p.text = "Tedious Document Verification:"
    p.level = 1
    p = tf2.add_paragraph()
    p.text = "Manual checks for driving licenses cause overhead and delays."
    p.level = 2
    p = tf2.add_paragraph()
    p.text = "Lack of Premium Options & Trust:"
    p.level = 1
    p = tf2.add_paragraph()
    p.text = "Difficult to locate specialized or premium vehicles securely."
    p.level = 2
    p = tf2.add_paragraph()
    p.text = "Poor User Interfaces:"
    p.level = 1
    p = tf2.add_paragraph()
    p.text = "Existing solutions are often unintuitive and do not cater properly to both owners and renters."
    p.level = 2

    # --- Slide 3: What Are We Building? (The Solution) ---
    slide3 = prs.slides.add_slide(bullet_slide_layout)
    shapes3 = slide3.shapes
    title3 = shapes3.title
    body3 = shapes3.placeholders[1]
    title3.text = "What Are We Building? (The Solution)"
    tf3 = body3.text_frame
    
    p = tf3.paragraphs[0]
    p.text = "A comprehensive Premium Vehicle Rental Web Application."
    p.level = 0
    p = tf3.add_paragraph()
    p.text = "Connects vehicle owners with users looking to rent premium vehicles."
    p.level = 1
    p = tf3.add_paragraph()
    p.text = "It eliminates traditional hurdles by providing a single, unified digital platform."
    p.level = 1
    p = tf3.add_paragraph()
    p.text = "Key Components:"
    p.level = 0
    p = tf3.add_paragraph()
    p.text = "Renter Module: For browsing, filtering, and seamless one-click booking."
    p.level = 1
    p = tf3.add_paragraph()
    p.text = "Owner Module: For securely listing idle vehicles to generate income."
    p.level = 1
    p = tf3.add_paragraph()
    p.text = "Admin Module: For governing the platform, moderating users, and managing operations."
    p.level = 1

    # --- Slide 4: Target Audience ---
    slide4 = prs.slides.add_slide(bullet_slide_layout)
    shapes4 = slide4.shapes
    title4 = shapes4.title
    body4 = shapes4.placeholders[1]
    title4.text = "Target Audience"
    tf4 = body4.text_frame
    
    p = tf4.paragraphs[0]
    p.text = "Primary Renters:"
    p.level = 0
    p = tf4.add_paragraph()
    p.text = "Tourists/Travelers needing reliable, premium vehicles for vacations."
    p.level = 1
    p = tf4.add_paragraph()
    p.text = "Corporate clients traveling for business trips."
    p.level = 1
    p = tf4.add_paragraph()
    p.text = "Individuals needing high-end vehicles for special occasions (weddings, events)."
    p.level = 1
    p = tf4.add_paragraph()
    p.text = "Primary Owners:"
    p.level = 0
    p = tf4.add_paragraph()
    p.text = "People with premium vehicles seeking passive asset monetization."
    p.level = 1
    p = tf4.add_paragraph()
    p.text = "Small-scale fleet managers looking for a modern platform to list inventory."
    p.level = 1

    # --- Slide 5: Core Features ---
    slide5 = prs.slides.add_slide(bullet_slide_layout)
    shapes5 = slide5.shapes
    title5 = shapes5.title
    body5 = shapes5.placeholders[1]
    title5.text = "Core Features"
    tf5 = body5.text_frame
    
    p = tf5.paragraphs[0]
    p.text = "Smart Authentication & Security:"
    p.level = 0
    p = tf5.add_paragraph()
    p.text = "Secure JWT-based login and mandatory driver's license ID uploads."
    p.level = 1
    p = tf5.add_paragraph()
    p.text = "Dynamic Search Engine:"
    p.level = 0
    p = tf5.add_paragraph()
    p.text = "Filter vehicles by Real-time availability, Make, Model, and Price range."
    p.level = 1
    p = tf5.add_paragraph()
    p.text = "Scalable Owner Dashboards:"
    p.level = 0
    p = tf5.add_paragraph()
    p.text = "Owners can effortlessly add new listings, toggle availability, and track usage."
    p.level = 1
    p = tf5.add_paragraph()
    p.text = "Administrative Oversight:"
    p.level = 0
    p = tf5.add_paragraph()
    p.text = "Centralized view of all platform users for data integrity and moderation."
    p.level = 1

    # --- Slide 6: How This Project Is Created (Tech Stack) ---
    slide6 = prs.slides.add_slide(bullet_slide_layout)
    shapes6 = slide6.shapes
    title6 = shapes6.title
    body6 = shapes6.placeholders[1]
    title6.text = "How This Project Is Created (Tech Stack)"
    tf6 = body6.text_frame
    
    p = tf6.paragraphs[0]
    p.text = "Frontend Development (Client-Side):"
    p.level = 0
    p = tf6.add_paragraph()
    p.text = "HTML5, CSS3, and Vanilla JavaScript (ES6+)."
    p.level = 1
    p = tf6.add_paragraph()
    p.text = "Custom state-management system (store.js) optimized for Single Page Application routing."
    p.level = 1
    p = tf6.add_paragraph()
    p.text = "Backend Component (Server-Side):"
    p.level = 0
    p = tf6.add_paragraph()
    p.text = "Python with the Flask framework mapping modular RESTful API endpoints."
    p.level = 1
    p = tf6.add_paragraph()
    p.text = "Database Environment:"
    p.level = 0
    p = tf6.add_paragraph()
    p.text = "SQLite, providing a zero-configuration, transactional SQL database engine."
    p.level = 1

    # --- Slide 7: System Architecture & Workflow ---
    slide7 = prs.slides.add_slide(bullet_slide_layout)
    shapes7 = slide7.shapes
    title7 = shapes7.title
    body7 = shapes7.placeholders[1]
    title7.text = "System Architecture & Workflow"
    tf7 = body7.text_frame
    
    p = tf7.paragraphs[0]
    p.text = "Client Application Structure:"
    p.level = 0
    p = tf7.add_paragraph()
    p.text = "Component-based UI modules corresponding to separate Javascript Views (e.g. HomeView, ExploriotView)."
    p.level = 1
    p = tf7.add_paragraph()
    p.text = "Centralized Store Pattern:"
    p.level = 0
    p = tf7.add_paragraph()
    p.text = "A universal JS 'store' governs authentication tokens and real-time user state (avoiding prop drilling)."
    p.level = 1
    p = tf7.add_paragraph()
    p.text = "Backend Micro-Architecture:"
    p.level = 0
    p = tf7.add_paragraph()
    p.text = "Separation of concerns using auth.py (Handling users) and routes to models.py (Handling DB Queries)."
    p.level = 1
    p = tf7.add_paragraph()
    p.text = "Request Lifecycle: Browser Event -> Store Dispatch -> API Fetch Endpoint -> DB Transaction & JSON Response."
    p.level = 1

    # --- Slide 8: Applications & Use Cases ---
    slide8 = prs.slides.add_slide(bullet_slide_layout)
    shapes8 = slide8.shapes
    title8 = shapes8.title
    body8 = shapes8.placeholders[1]
    title8.text = "Applications & Use Cases"
    tf8 = body8.text_frame
    
    p = tf8.paragraphs[0]
    p.text = "Use Case 1: The Urgent Renter"
    p.level = 0
    p = tf8.add_paragraph()
    p.text = "A user arrives locally and needs a high-end car tomorrow. They register, upload their license, and book a BMW instantly online with guaranteed approval via the platform."
    p.level = 1
    p = tf8.add_paragraph()
    p.text = "Use Case 2: The Idle Asset Owner"
    p.level = 0
    p = tf8.add_paragraph()
    p.text = "An owner with an unused luxury SUV signs up as an owner, lists vehicle details, sets the price, and instantly exposes their car to verified renters."
    p.level = 1
    p = tf8.add_paragraph()
    p.text = "Use Case 3: The Administrator"
    p.level = 0
    p = tf8.add_paragraph()
    p.text = "An Admin logs in to verify platform stats, viewing all current users and ensuring all owner listings comply with platform standards."
    p.level = 1

    # --- Slide 9: Project Progress (Completed vs. Left) ---
    slide9 = prs.slides.add_slide(bullet_slide_layout)
    shapes9 = slide9.shapes
    title9 = shapes9.title
    body9 = shapes9.placeholders[1]
    title9.text = "Project Progress"
    tf9 = body9.text_frame
    
    p = tf9.paragraphs[0]
    p.text = "Successfully Completed Elements:"
    p.level = 0
    p = tf9.add_paragraph()
    p.text = "Routing mechanism and entirely modularized UI/UX."
    p.level = 1
    p = tf9.add_paragraph()
    p.text = "User Auth (Register/Login via secure JWT tokens), including Admin scopes."
    p.level = 1
    p = tf9.add_paragraph()
    p.text = "Sophisticated Database initialization with a seeded roster of Indian premium vehicles."
    p.level = 1
    p = tf9.add_paragraph()
    p.text = "Admin dashboard and Owner listing workflows operational."
    p.level = 1
    
    p = tf9.add_paragraph()
    p.text = "Pending / Left to Implement:"
    p.level = 0
    p = tf9.add_paragraph()
    p.text = "Complete Integration of Live Payment Gateways (such as Stripe)."
    p.level = 1
    p = tf9.add_paragraph()
    p.text = "Integration of external email/SMS notification APIs for immediate reservation alerts."
    p.level = 1

    # --- Slide 10: Upcoming Plans to Enhance It ---
    slide10 = prs.slides.add_slide(bullet_slide_layout)
    shapes10 = slide10.shapes
    title10 = shapes10.title
    body10 = shapes10.placeholders[1]
    title10.text = "Upcoming Plans / Enhancements"
    tf10 = body10.text_frame
    
    p = tf10.paragraphs[0]
    p.text = "Native Mobile Integration:"
    p.level = 0
    p = tf10.add_paragraph()
    p.text = "Migrating the platform logic to React Native or Flutter to spin off dedicated iOS and Android apps."
    p.level = 1
    p = tf10.add_paragraph()
    p.text = "AI-Driven Personalization:"
    p.level = 0
    p = tf10.add_paragraph()
    p.text = "Introducing machine learning to recommend vehicles based on past user rentals and seasonal trends."
    p.level = 1
    p = tf10.add_paragraph()
    p.text = "Dynamic Pricing Intelligence:"
    p.level = 0
    p = tf10.add_paragraph()
    p.text = "Developing algorithms to raise or lower vehicle prices autonomously depending on holiday demand."
    p.level = 1
    p = tf10.add_paragraph()
    p.text = "Eco-Friendly Vertical:"
    p.level = 0
    p = tf10.add_paragraph()
    p.text = "Expanding our inventory to heavily feature and promote Premium Electric Vehicles (EVs)."
    p.level = 1

    # --- Slide 11: Conclusion ---
    slide11 = prs.slides.add_slide(title_slide_layout)
    title11 = slide11.shapes.title
    subtitle11 = slide11.placeholders[1]
    title11.text = "Thank You!"
    subtitle11.text = "I am open to any questions.\n\n[Your Contact / Github Info]"

    # Save the presentation
    output_path = os.path.join(os.getcwd(), 'RentWheels_MCA_Presentation.pptx')
    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == '__main__':
    create_presentation()
