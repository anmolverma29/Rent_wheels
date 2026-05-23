import os
import datetime
try:
    from docx import Document
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    print("Please install python-docx using: pip install python-docx")
    exit(1)

def create_word_doc():
    document = Document()

    # Document Title / Slide 1
    head = document.add_heading('RentWheels', 0)
    head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    date_str = datetime.date.today().strftime("%B %d, %Y")
    subtitle = document.add_paragraph(f"MCA Final Year Project Documentation\nPremium Vehicle Rental Platform\n\n[Your Name] | {date_str}")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    document.add_page_break()

    # Slide 2: Problem Statement
    document.add_heading('The Problem Statement', level=1)
    document.add_paragraph("Current vehicle rental processes are highly fragmented and outdated.")
    
    document.add_paragraph("Tedious Document Verification:", style='List Bullet')
    document.add_paragraph("Manual checks for driving licenses cause overhead and delays.", style='List Bullet 2')
    
    document.add_paragraph("Lack of Premium Options & Trust:", style='List Bullet')
    document.add_paragraph("Difficult to locate specialized or premium vehicles securely.", style='List Bullet 2')
    
    document.add_paragraph("Poor User Interfaces:", style='List Bullet')
    document.add_paragraph("Existing solutions are often unintuitive and do not cater properly to both owners and renters.", style='List Bullet 2')

    # Slide 3: What Are We Building? (The Solution)
    document.add_heading('What Are We Building? (The Solution)', level=1)
    document.add_paragraph("A comprehensive Premium Vehicle Rental Web Application.", style='List Bullet')
    document.add_paragraph("Connects vehicle owners with users looking to rent premium vehicles.", style='List Bullet')
    document.add_paragraph("It eliminates traditional hurdles by providing a single, unified digital platform.", style='List Bullet')
    
    document.add_paragraph("Key Components:", style='List Bullet')
    document.add_paragraph("Renter Module: For browsing, filtering, and seamless one-click booking.", style='List Bullet 2')
    document.add_paragraph("Owner Module: For securely listing idle vehicles to generate income.", style='List Bullet 2')
    document.add_paragraph("Admin Module: For governing the platform, moderating users, and managing operations.", style='List Bullet 2')

    # Slide 4: Target Audience
    document.add_heading('Target Audience', level=1)
    document.add_paragraph("Primary Renters:", style='List Bullet')
    document.add_paragraph("Tourists/Travelers needing reliable, premium vehicles for vacations.", style='List Bullet 2')
    document.add_paragraph("Corporate clients traveling for business trips.", style='List Bullet 2')
    document.add_paragraph("Individuals needing high-end vehicles for special occasions (weddings, events).", style='List Bullet 2')
    
    document.add_paragraph("Primary Owners:", style='List Bullet')
    document.add_paragraph("People with premium vehicles seeking passive asset monetization.", style='List Bullet 2')
    document.add_paragraph("Small-scale fleet managers looking for a modern platform to list inventory.", style='List Bullet 2')

    # Slide 5: Core Features
    document.add_heading('Core Features', level=1)
    document.add_paragraph("Smart Authentication & Security:", style='List Bullet')
    document.add_paragraph("Secure JWT-based login and mandatory driver's license ID uploads.", style='List Bullet 2')
    
    document.add_paragraph("Dynamic Search Engine:", style='List Bullet')
    document.add_paragraph("Filter vehicles by Real-time availability, Make, Model, and Price range.", style='List Bullet 2')
    
    document.add_paragraph("Scalable Owner Dashboards:", style='List Bullet')
    document.add_paragraph("Owners can effortlessly add new listings, toggle availability, and track usage.", style='List Bullet 2')
    
    document.add_paragraph("Administrative Oversight:", style='List Bullet')
    document.add_paragraph("Centralized view of all platform users for data integrity and moderation.", style='List Bullet 2')

    # Slide 6: How This Project Is Created (Tech Stack)
    document.add_heading('How This Project Is Created (Tech Stack)', level=1)
    document.add_paragraph("Frontend Development (Client-Side):", style='List Bullet')
    document.add_paragraph("HTML5, CSS3, and Vanilla JavaScript (ES6+).", style='List Bullet 2')
    document.add_paragraph("Custom state-management system (store.js) optimized for Single Page Application routing.", style='List Bullet 2')
    
    document.add_paragraph("Backend Component (Server-Side):", style='List Bullet')
    document.add_paragraph("Python with the Flask framework mapping modular RESTful API endpoints.", style='List Bullet 2')
    
    document.add_paragraph("Database Environment:", style='List Bullet')
    document.add_paragraph("SQLite, providing a zero-configuration, transactional SQL database engine.", style='List Bullet 2')

    # Slide 7: System Architecture & Workflow
    document.add_heading('System Architecture & Workflow', level=1)
    document.add_paragraph("Client Application Structure:", style='List Bullet')
    document.add_paragraph("Component-based UI modules corresponding to separate Javascript Views (e.g. HomeView, ExploriotView).", style='List Bullet 2')
    
    document.add_paragraph("Centralized Store Pattern:", style='List Bullet')
    document.add_paragraph("A universal JS 'store' governs authentication tokens and real-time user state (avoiding prop drilling).", style='List Bullet 2')
    
    document.add_paragraph("Backend Micro-Architecture:", style='List Bullet')
    document.add_paragraph("Separation of concerns using auth.py (Handling users) and routes to models.py (Handling DB Queries).", style='List Bullet 2')
    
    document.add_paragraph("Request Lifecycle:", style='List Bullet')
    document.add_paragraph("Browser Event -> Store Dispatch -> API Fetch Endpoint -> DB Transaction & JSON Response.", style='List Bullet 2')

    # Slide 8: Applications & Use Cases
    document.add_heading('Applications & Use Cases', level=1)
    document.add_paragraph("Use Case 1: The Urgent Renter", style='List Bullet')
    document.add_paragraph("A user arrives locally and needs a high-end car tomorrow. They register, upload their license, and book a BMW instantly online with guaranteed approval via the platform.", style='List Bullet 2')
    
    document.add_paragraph("Use Case 2: The Idle Asset Owner", style='List Bullet')
    document.add_paragraph("An owner with an unused luxury SUV signs up as an owner, lists vehicle details, sets the price, and instantly exposes their car to verified renters.", style='List Bullet 2')
    
    document.add_paragraph("Use Case 3: The Administrator", style='List Bullet')
    document.add_paragraph("An Admin logs in to verify platform stats, viewing all current users and ensuring all owner listings comply with platform standards.", style='List Bullet 2')

    # Slide 9: Project Progress (Completed vs. Left)
    document.add_heading('Project Progress (Completed vs. Left)', level=1)
    document.add_paragraph("Successfully Completed Elements:", style='List Bullet')
    document.add_paragraph("Routing mechanism and entirely modularized UI/UX.", style='List Bullet 2')
    document.add_paragraph("User Auth (Register/Login via secure JWT tokens), including Admin scopes.", style='List Bullet 2')
    document.add_paragraph("Sophisticated Database initialization with a seeded roster of Indian premium vehicles.", style='List Bullet 2')
    document.add_paragraph("Admin dashboard and Owner listing workflows operational.", style='List Bullet 2')
    
    document.add_paragraph("Pending / Left to Implement:", style='List Bullet')
    document.add_paragraph("Complete Integration of Live Payment Gateways (such as Stripe).", style='List Bullet 2')
    document.add_paragraph("Integration of external email/SMS notification APIs for immediate reservation alerts.", style='List Bullet 2')

    # Slide 10: Upcoming Plans to Enhance It
    document.add_heading('Upcoming Plans / Enhancements', level=1)
    document.add_paragraph("Native Mobile Integration:", style='List Bullet')
    document.add_paragraph("Migrating the platform logic to React Native or Flutter to spin off dedicated iOS and Android apps.", style='List Bullet 2')
    
    document.add_paragraph("AI-Driven Personalization:", style='List Bullet')
    document.add_paragraph("Introducing machine learning to recommend vehicles based on past user rentals and seasonal trends.", style='List Bullet 2')
    
    document.add_paragraph("Dynamic Pricing Intelligence:", style='List Bullet')
    document.add_paragraph("Developing algorithms to raise or lower vehicle prices autonomously depending on holiday demand.", style='List Bullet 2')
    
    document.add_paragraph("Eco-Friendly Vertical:", style='List Bullet')
    document.add_paragraph("Expanding our inventory to heavily feature and promote Premium Electric Vehicles (EVs).", style='List Bullet 2')

    # Save the document
    output_path = os.path.join(os.getcwd(), 'RentWheels_MCA_Documentation.docx')
    document.save(output_path)
    print(f"Word document saved successfully to: {output_path}")

if __name__ == '__main__':
    create_word_doc()
